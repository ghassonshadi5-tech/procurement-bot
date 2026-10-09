"""
Company Owner - Reports and Excel export handlers.
Aligned with Procurement Bot Architecture v3.0 and Schema v2.0.
"""
from typing import List, Optional
from datetime import datetime
from io import BytesIO

from aiogram import Router, F
from aiogram.types import Message, CallbackQuery, InlineKeyboardMarkup, InlineKeyboardButton
from sqlalchemy import select, func

from app.core.logger import get_logger
from app.database.database import get_session
from app.database.models.purchase_request import PurchaseRequest
from app.database.models.rfq import Rfq
from app.database.models.supplier_relation import SupplierRelation

from app.bot.handlers.company_owner._helpers import (
    IsBuyerSide,
    require_buyer_admin_access,
)

logger = get_logger(__name__)
router = Router(name="company_owner_reports")


# ============================================================================
# Button: "📊 التقارير"
# ============================================================================
@router.message(F.text == "📊 التقارير", IsBuyerSide())
async def btn_company_reports(
    message: Message,
    user,
    roles: List[str],
    company_id: Optional[int]
):
    """Handle '📊 التقارير' button."""
    logger.info(
        "HANDLER_HIT | btn_company_reports | tg_id=%s | roles=%s",
        message.from_user.id,
        roles,
    )

    allowed_roles = ["company_owner", "manager", "accountant"]
    has_permission = any(role in allowed_roles for role in roles)

    if not has_permission:
        logger.warning(
            "ACCESS_DENIED | reports | tg_id=%s | roles=%s",
            message.from_user.id,
            roles,
            extra={"security": True},
        )
        await message.answer("⛔ غير مصرح لك بالوصول لهذه الوظيفة.")
        return

    if not user or not company_id:
        await message.answer("⚠️ لم يتم العثور على شركتك.")
        return

    try:
        async with get_session(company_id=company_id) as session:
            stmt_pr = (
                select(PurchaseRequest.status, func.count(PurchaseRequest.id))
                .where(PurchaseRequest.company_id == company_id)
                .group_by(PurchaseRequest.status)
            )
            result_pr = await session.execute(stmt_pr)
            pr_stats = result_pr.all()

            stmt_rfq = (
                select(Rfq.status, func.count(Rfq.id))
                .where(Rfq.company_id == company_id)
                .group_by(Rfq.status)
            )
            result_rfq = await session.execute(stmt_rfq)
            rfq_stats = result_rfq.all()

            stmt_suppliers = (
                select(func.count(SupplierRelation.id))
                .where(
                    SupplierRelation.company_id == company_id,
                    SupplierRelation.status == "active"
                )
            )
            result_suppliers = await session.execute(stmt_suppliers)
            active_suppliers = result_suppliers.scalar() or 0

        response = (
            f"📊 <b>تقارير الشركة</b>\n"
            f"━━━━━━━━━━━━━━━━━━━\n\n"
            f"🛒 <b>طلبات الشراء:</b>\n"
        )

        for status, count in pr_stats:
            status_val = status.value if hasattr(status, 'value') else str(status)
            response += f"   • {status_val}: <b>{count}</b>\n"

        response += f"\n📨 <b>طلبات عروض الأسعار (RFQ):</b>\n"

        for status, count in rfq_stats:
            status_val = status.value if hasattr(status, 'value') else str(status)
            response += f"   • {status_val}: <b>{count}</b>\n"

        response += (
            f"\n🏪 <b>الموردون النشطون:</b> {active_suppliers}\n\n"
            f"━━━━━━━━━━━━━━━━━━━\n"
            f"<i>📅 تم إنشاء التقرير في: {datetime.now().strftime('%Y-%m-%d %H:%M')}</i>"
        )

        keyboard = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(
                text="📥 تصدير إلى Excel",
                callback_data="co_export_report_excel"
            )]
        ])

        await message.answer(response, reply_markup=keyboard, parse_mode="HTML")

    except Exception:
        logger.exception("Error generating company report")
        await message.answer("❌ عذراً، حدث خطأ أثناء إنشاء التقرير.")


# ============================================================================
# Callback: Export Report to Excel
# ============================================================================
@router.callback_query(F.data == "co_export_report_excel")
async def callback_export_report(
    callback: CallbackQuery,
    user,
    roles: List[str],
    company_id: Optional[int]
):
    """Handle export report button - Generates actual Excel file."""
    if not require_buyer_admin_access(roles):
        await callback.answer("⛔ غير مصرح لك.", show_alert=True)
        return

    if not company_id:
        await callback.answer("⚠️ لم يتم العثور على شركتك.", show_alert=True)
        return

    try:
        from app.database.models.supplier import Supplier
        from aiogram.types import BufferedInputFile
        import openpyxl

        async with get_session(company_id=company_id) as session:
            pr_stmt = (
                select(PurchaseRequest)
                .where(PurchaseRequest.company_id == company_id)
                .order_by(PurchaseRequest.created_at.desc())
                .limit(100)
            )
            pr_res = await session.execute(pr_stmt)
            requests = pr_res.scalars().all()

            rfq_stmt = (
                select(Rfq)
                .where(Rfq.company_id == company_id)
                .order_by(Rfq.created_at.desc())
                .limit(100)
            )
            rfq_res = await session.execute(rfq_stmt)
            rfqs = rfq_res.scalars().all()

            sup_stmt = (
                select(SupplierRelation, Supplier)
                .join(Supplier, SupplierRelation.supplier_id == Supplier.id)
                .where(
                    SupplierRelation.company_id == company_id,
                    SupplierRelation.status == "active",
                )
            )
            sup_res = await session.execute(sup_stmt)
            suppliers = sup_res.all()

        wb = openpyxl.Workbook()

        # Sheet 1: Purchase Requests
        ws_pr = wb.active
        ws_pr.title = "طلبات الشراء"
        ws_pr.append(["رقم الطلب", "العنوان", "الحالة", "تاريخ الإنشاء", "مقدم الطلب"])
        for req in requests:
            status_val = req.status.value if hasattr(req.status, "value") else str(req.status)
            ws_pr.append([
                req.id,
                req.title,
                status_val,
                req.created_at.strftime("%Y-%m-%d"),
                "-",
            ])

        # Sheet 2: RFQs
        ws_rfq = wb.create_sheet("طلبات العروض")
        ws_rfq.append(["رقم RFQ", "الأسلوب", "الحالة", "الموعد النهائي", "تاريخ الإنشاء"])
        for rfq in rfqs:
            status_val = rfq.status.value if hasattr(rfq.status, "value") else str(rfq.status)
            method_val = rfq.method.value if hasattr(rfq.method, "value") else str(rfq.method)
            dl = rfq.deadline.strftime("%Y-%m-%d") if rfq.deadline else "-"
            ws_rfq.append([
                rfq.id,
                method_val,
                status_val,
                dl,
                rfq.created_at.strftime("%Y-%m-%d"),
            ])

        # Sheet 3: Suppliers
        ws_sup = wb.create_sheet("الموردون")
        ws_sup.append(["اسم المورد", "النشاط", "المدن", "الحالة"])
        for rel, sup in suppliers:
            cities = ", ".join(sup.cities) if sup.cities else "-"
            status_val = rel.status.value if hasattr(rel.status, "value") else str(rel.status)
            ws_sup.append([sup.name, sup.activity or "-", cities, status_val])

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        filename = f"report_co_{company_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        doc = BufferedInputFile(buffer.read(), filename=filename)

        await callback.message.answer_document(
            doc,
            caption="📥 <b>تقرير الشركة (Excel)</b>",
            parse_mode="HTML",
        )
        await callback.answer("✅ تم التصدير بنجاح", show_alert=False)

    except ImportError:
        logger.error("openpyxl not installed.")
        await callback.answer("❌ مكتبة Excel غير مثبتة على الخادم.", show_alert=True)
    except Exception:
        logger.exception("Excel export error")
        await callback.answer("❌ فشل التصدير. حاول لاحقاً.", show_alert=True)