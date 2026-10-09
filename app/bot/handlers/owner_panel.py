"""
Super Admin Panel handlers - admin panel for the platform owner.
Aligned with Procurement Bot Architecture v3.0 and Schema v2.0.

Handles:
- Company registration requests (approval/rejection).
- Supplier registration requests (approval/rejection).
- Company management and Supplier management.
- System-wide statistics and Excel export.
- Platform settings display.

KEY FIXES:
- Added handler for "🏭 إدارة الموردين" (was missing).
- Fixed button text mismatches (Reports and Settings).
- Fixed Telegram ID vs Internal User ID in approve/reject callbacks.
- Aligned all handler texts with get_super_admin_keyboard() in main_menu.py.

IMPROVEMENTS (this revision):
- Uses get_logger(__name__) → correlation IDs in all logs.
- Unscoped sessions use silent=True (platform-level operations).
- Passes the rejection reason to supplier_service.reject_supplier()
  so it is recorded in audit_log.
- Marks sensitive log lines with extra={"security": True}.
- Logging uses %s formatting.
"""
import io
from datetime import datetime, timezone, timedelta
from typing import List, Optional

from aiogram import Router, F
from aiogram.filters import Command
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import Message, CallbackQuery, InlineKeyboardMarkup, InlineKeyboardButton, BufferedInputFile
from sqlalchemy import select, func
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.core.logger import get_logger
from app.config.settings import settings
from app.core import permissions
from app.database.database import get_session
from app.database.models.company import Company, CompanyStatus
from app.database.models.supplier import Supplier
from app.database.models.user import User, UserStatus
from app.database.models.purchase_request import PurchaseRequest
from app.services.user_service import user_service
from app.services.company_service import company_service
from app.services.supplier_service import supplier_service


logger = get_logger(__name__)
router = Router(name="owner_panel")


# ============================================================================
# FSM States for rejection reason
# ============================================================================
class OwnerRejectFlow(StatesGroup):
    """States for rejecting a company/supplier request with reason."""
    waiting_for_company_rejection_reason = State()
    waiting_for_supplier_rejection_reason = State()


# ============================================================================
# Helper: Check Super Admin Permission
# ============================================================================
def require_super_admin(roles: List[str]) -> bool:
    """Check if the user has super_admin role."""
    return permissions.is_super_admin(roles)


# ============================================================================
# Command: /admin (Shortcut to Super Admin Panel)
# ============================================================================
@router.message(Command("admin"))
async def cmd_admin(message: Message, roles: List[str]):
    """Handle /admin command - show Super Admin dashboard."""
    if not require_super_admin(roles):
        await message.answer("⛔ غير مصرح لك بالوصول لهذه الوظيفة.")
        return

    from app.bot.keyboards.main_menu import get_super_admin_keyboard
    kb = get_super_admin_keyboard()
    await message.answer(
        "👑 <b>لوحة التحكم العامة (Super Admin)</b>",
        reply_markup=kb,
        parse_mode="HTML",
    )


# ============================================================================
# Button: "📨 طلبات تسجيل الشركات"
# ============================================================================
@router.message(F.text == "📨 طلبات تسجيل الشركات")
async def btn_company_requests(message: Message, roles: List[str]):
    """Handle the '📨 طلبات تسجيل الشركات' button."""
    if not require_super_admin(roles):
        await message.answer("⛔ غير مصرح لك بالوصول لهذه الوظيفة.")
        return

    logger.info("Super Admin %s viewing company requests", message.from_user.id)

    try:
        pending_companies = await company_service.list_all_companies(
            status=CompanyStatus.pending
        )

        if not pending_companies:
            await message.answer(
                "✅ لا توجد طلبات تسجيل شركات معلقة حالياً.\n"
                "كل الطلبات تمت معالجتها! 🎉"
            )
            return

        await message.answer(
            f"📨 <b>طلبات تسجيل الشركات المعلقة ({len(pending_companies)} طلب):</b>",
            parse_mode="HTML",
        )

        for i, company in enumerate(pending_companies[:10], 1):
            date_str = (
                company.created_at.strftime("%Y-%m-%d %H:%M")
                if company.created_at
                else "N/A"
            )

            # Fetch the user who created the request
            creator_name = "غير معروف"
            creator_tg_id = "N/A"
            if company.created_by:
                async with get_session(silent=True) as session:
                    stmt = select(User).where(User.id == company.created_by)
                    result = await session.execute(stmt)
                    creator = result.scalar_one_or_none()
                    if creator:
                        creator_name = creator.full_name
                        creator_tg_id = creator.telegram_user_id

            keyboard = InlineKeyboardMarkup(inline_keyboard=[
                [
                    InlineKeyboardButton(
                        text="✅ اعتماد",
                        callback_data=f"sa_approve_company:{company.id}",
                    ),
                    InlineKeyboardButton(
                        text="❌ رفض مع سبب",
                        callback_data=f"sa_reject_company_start:{company.id}",
                    ),
                ]
            ])

            card_text = (
                f"<b>{i}. 🏢 {company.name}</b>\n"
                f"👤 مقدم الطلب: {creator_name}\n"
                f"🆔 Telegram ID: <code>{creator_tg_id}</code>\n"
                f"🏙️ المدينة: {company.city or 'غير محدد'}\n"
                f"📅 التاريخ: {date_str}\n"
                f"🔢 رقم الشركة: <code>{company.id}</code>"
            )

            await message.answer(card_text, reply_markup=keyboard, parse_mode="HTML")

        if len(pending_companies) > 10:
            await message.answer(
                f"<i>... و {len(pending_companies) - 10} طلبات أخرى</i>",
                parse_mode="HTML",
            )

    except Exception:
        logger.exception("Error fetching company requests")
        await message.answer("❌ عذراً، حدث خطأ أثناء جلب الطلبات.")


# ============================================================================
# Button: "📨 طلبات تسجيل الموردين"
# ============================================================================
@router.message(F.text == "📨 طلبات تسجيل الموردين")
async def btn_supplier_requests(message: Message, roles: List[str]):
    """Handle the '📨 طلبات تسجيل الموردين' button."""
    if not require_super_admin(roles):
        await message.answer("⛔ غير مصرح لك بالوصول لهذه الوظيفة.")
        return

    try:
        pending_suppliers = await supplier_service.list_all_suppliers(
            status=CompanyStatus.pending
        )

        if not pending_suppliers:
            await message.answer(
                "✅ لا توجد طلبات تسجيل موردين معلقة حالياً.\n"
                "كل الطلبات تمت معالجتها! 🎉"
            )
            return

        await message.answer(
            f"📨 <b>طلبات تسجيل الموردين المعلقة ({len(pending_suppliers)} طلب):</b>",
            parse_mode="HTML",
        )

        for i, supplier in enumerate(pending_suppliers[:10], 1):
            date_str = (
                supplier.created_at.strftime("%Y-%m-%d %H:%M")
                if supplier.created_at
                else "N/A"
            )

            keyboard = InlineKeyboardMarkup(inline_keyboard=[
                [
                    InlineKeyboardButton(
                        text="✅ اعتماد",
                        callback_data=f"sa_approve_supplier:{supplier.id}",
                    ),
                    InlineKeyboardButton(
                        text="❌ رفض",
                        callback_data=f"sa_reject_supplier_start:{supplier.id}",
                    ),
                ]
            ])

            card_text = (
                f"<b>{i}. 🏭 {supplier.name}</b>\n"
                f"🏙️ المدن: {', '.join(supplier.cities) if supplier.cities else 'غير محدد'}\n"
                f"📅 التاريخ: {date_str}\n"
                f"🔢 رقم المورد: <code>{supplier.id}</code>"
            )

            await message.answer(card_text, reply_markup=keyboard, parse_mode="HTML")

    except Exception:
        logger.exception("Error fetching supplier requests")
        await message.answer("❌ عذراً، حدث خطأ أثناء جلب الطلبات.")


# ============================================================================
# Button: "🏢 إدارة الشركات"
# ============================================================================
@router.message(F.text == "🏢 إدارة الشركات")
async def btn_manage_companies(message: Message, roles: List[str]):
    """Handle the '🏢 إدارة الشركات' button."""
    if not require_super_admin(roles):
        await message.answer("⛔ غير مصرح لك بالوصول لهذه الوظيفة.")
        return

    try:
        active_companies = await company_service.list_all_companies(
            status=CompanyStatus.active
        )

        if not active_companies:
            await message.answer("📭 لا توجد شركات نشطة حالياً.")
            return

        response = f"🏢 <b>الشركات النشطة ({len(active_companies)} شركة):</b>\n"

        for i, company in enumerate(active_companies[:15], 1):
            date_str = (
                company.created_at.strftime("%Y-%m-%d")
                if company.created_at
                else "N/A"
            )

            # Count users for this company
            async with get_session(silent=True) as session:
                stmt = select(func.count(User.id)).where(
                    User.company_id == company.id,
                    User.status == UserStatus.active,
                )
                result = await session.execute(stmt)
                user_count = result.scalar() or 0

            response += (
                f"\n<b>{i}. 🏢 {company.name}</b>\n"
                f"   👥 المستخدمون: {user_count}\n"
                f"   🏙️ المدينة: {company.city or 'غير محدد'}\n"
                f"   📅 التسجيل: {date_str}"
            )

        if len(active_companies) > 15:
            response += f"\n\n<i>... و {len(active_companies) - 15} شركات أخرى</i>"

        await message.answer(response, parse_mode="HTML")

    except Exception:
        logger.exception("Error fetching companies")
        await message.answer("❌ عذراً، حدث خطأ أثناء جلب الشركات.")


# ============================================================================
# Button: "🏭 إدارة الموردين"
# ============================================================================
@router.message(F.text == "🏭 إدارة الموردين")
async def btn_manage_suppliers(message: Message, roles: List[str]):
    """Handle the '🏭 إدارة الموردين' button - list all active suppliers."""
    if not require_super_admin(roles):
        await message.answer("⛔ غير مصرح لك بالوصول لهذه الوظيفة.")
        return

    try:
        active_suppliers = await supplier_service.list_all_suppliers(
            status=CompanyStatus.active
        )

        if not active_suppliers:
            await message.answer("📭 لا يوجد موردون نشطون حالياً.")
            return

        response = f"🏭 <b>الموردون النشطون ({len(active_suppliers)} مورد):</b>\n"

        for i, supplier in enumerate(active_suppliers[:15], 1):
            date_str = (
                supplier.created_at.strftime("%Y-%m-%d")
                if supplier.created_at
                else "N/A"
            )
            cities = ", ".join(supplier.cities) if supplier.cities else "غير محدد"
            is_verified = "✅" if supplier.is_verified else "⏳"

            # Count items in supplier catalog
            async with get_session(silent=True) as session:
                from app.database.models.item import Item
                stmt = select(func.count(Item.id)).where(
                    Item.supplier_id == supplier.id,
                    Item.is_active.is_(True),
                )
                result = await session.execute(stmt)
                item_count = result.scalar() or 0

            response += (
                f"\n<b>{i}. 🏭 {supplier.name}</b>\n"
                f"   📦 الأصناف في الكتالوج: {item_count}\n"
                f"   🏙️ المدن: {cities}\n"
                f"   🏷️ النشاط: {supplier.activity or 'غير محدد'}\n"
                f"   ✅ التوثيق: {is_verified}\n"
                f"   📅 التسجيل: {date_str}"
            )

        if len(active_suppliers) > 15:
            response += f"\n\n<i>... و {len(active_suppliers) - 15} موردون آخرون</i>"

        await message.answer(response, parse_mode="HTML")

    except Exception:
        logger.exception("Error fetching suppliers")
        await message.answer("❌ عذراً، حدث خطأ أثناء جلب الموردين.")


# ============================================================================
# Button: "📊 التقارير العامة"
# ============================================================================
@router.message(F.text == "📊 التقارير العامة")
async def btn_reports(message: Message, roles: List[str]):
    """Handle the '📊 التقارير العامة' button."""
    if not require_super_admin(roles):
        await message.answer("⛔ غير مصرح لك بالوصول لهذه الوظيفة.")
        return

    try:
        async with get_session(silent=True) as session:
            total_companies = (await session.execute(
                select(func.count(Company.id)).where(
                    Company.status == CompanyStatus.active
                )
            )).scalar() or 0

            total_users = (await session.execute(
                select(func.count(User.id)).where(
                    User.status == UserStatus.active
                )
            )).scalar() or 0

            pending_companies = (await session.execute(
                select(func.count(Company.id)).where(
                    Company.status == CompanyStatus.pending
                )
            )).scalar() or 0

            total_suppliers = (await session.execute(
                select(func.count(Supplier.id)).where(
                    Supplier.status == CompanyStatus.active
                )
            )).scalar() or 0

            pending_suppliers = (await session.execute(
                select(func.count(Supplier.id)).where(
                    Supplier.status == CompanyStatus.pending
                )
            )).scalar() or 0

            total_purchase_requests = (await session.execute(
                select(func.count(PurchaseRequest.id))
            )).scalar() or 0

            week_ago = datetime.now(timezone.utc) - timedelta(days=7)
            new_companies_week = (await session.execute(
                select(func.count(Company.id)).where(
                    Company.created_at >= week_ago,
                    Company.status == CompanyStatus.active,
                )
            )).scalar() or 0

        report_text = (
            f"📊 <b>تقارير النظام العامة</b>\n"
            f"━━━━━━━━━━━━━━━━━━━\n\n"

            f"🏢 <b>الشركات المشترية:</b>\n"
            f"   • نشطة: <b>{total_companies}</b>\n"
            f"   • جديدة (آخر 7 أيام): <b>{new_companies_week}</b>\n"
            f"   • ⏳ معلقة: <b>{pending_companies}</b>\n\n"

            f"🏭 <b>الموردون:</b>\n"
            f"   • نشطون: <b>{total_suppliers}</b>\n"
            f"   • ⏳ معلقون: <b>{pending_suppliers}</b>\n\n"

            f"👥 <b>المستخدمون:</b>\n"
            f"   • إجمالي المستخدمين النشطين: <b>{total_users}</b>\n\n"

            f"🛒 <b>طلبات الشراء:</b>\n"
            f"   • إجمالي الطلبات: <b>{total_purchase_requests}</b>\n\n"

            f"━━━━━━━━━━━━━━━━━━━\n"
            f"<i>📅 تم إنشاء التقرير في: {datetime.now().strftime('%Y-%m-%d %H:%M')}</i>"
        )

        export_keyboard = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(
                text="📥 تصدير إلى Excel",
                callback_data="sa_export_report_excel",
            )]
        ])

        await message.answer(report_text, reply_markup=export_keyboard, parse_mode="HTML")

    except Exception:
        logger.exception("Error generating report")
        await message.answer("❌ عذراً، حدث خطأ أثناء إنشاء التقرير.")


# ============================================================================
# Button: "⚙️ إعدادات المنصة"
# ============================================================================
@router.message(F.text == "⚙️ إعدادات المنصة")
async def btn_settings(message: Message, roles: List[str]):
    """Handle the '⚙️ إعدادات المنصة' button."""
    if not require_super_admin(roles):
        await message.answer("⛔ غير مصرح لك بالوصول لهذه الوظيفة.")
        return

    settings_text = (
        f"⚙️ <b>إعدادات المنصة</b>\n"
        f"━━━━━━━━━━━━━━━━━━━\n\n"
        f"🤖 <b>اسم البوت:</b> @{settings.BOT_USERNAME}\n"
        f"👑 <b>معرف المدير الفائق:</b> <code>{settings.SUPER_ADMIN_TELEGRAM_ID}</code>\n"
        f"🌍 <b>البيئة:</b> {settings.APP_ENV}\n"
        f"🔧 <b>وضع التصحيح:</b> {'✅ مفعل' if settings.DEBUG else '❌ معطل'}\n"
        f"📊 <b>مستوى السجلات:</b> {settings.LOG_LEVEL}\n\n"
        f"━━━━━━━━━━━━━━━━━━━\n"
        f"💡 <i>لتعديل الإعدادات، عدّل ملف .env وأعد تشغيل البوت.</i>"
    )

    await message.answer(settings_text, parse_mode="HTML")


# ============================================================================
# Callback: Approve Company (With Strict Logging & Fallback)
# ============================================================================
@router.callback_query(F.data.startswith("sa_approve_company:"))
async def callback_approve_company(callback: CallbackQuery, roles: List[str]):
    """Handle approve company button click."""
    if not require_super_admin(roles):
        await callback.answer("⛔ غير مصرح لك.", show_alert=True)
        return

    company_id = int(callback.data.split(":", 1)[1])

    try:
        # 1. Get Internal User ID from Telegram ID
        admin_user = await user_service.get_user_by_telegram_id(callback.from_user.id)
        if not admin_user:
            await callback.answer("❌ خطأ: لم يتم العثور على حساب المدير.", show_alert=True)
            return

        admin_internal_id = admin_user.id

        # 2. Approve the company
        success = await company_service.approve_company(company_id, admin_internal_id)
        if not success:
            await callback.answer("❌ فشل الاعتماد أو الشركة غير موجودة.", show_alert=True)
            return

        # 3. Fetch company and creator to link them
        creator_tg_id = None
        async with get_session(silent=True) as session:
            stmt = select(Company).where(Company.id == company_id)
            result = await session.execute(stmt)
            company = result.scalar_one()

            creator_id = company.created_by

            if not creator_id:
                logger.warning(
                    "Company has NO creator_id — cannot assign owner role automatically"
                    " | company_id=%s",
                    company_id,
                    extra={"security": True},
                )
                await callback.message.edit_text(
                    f"✅ <b>تم اعتماد الشركة بنجاح!</b>\n"
                    f"🏢 الشركة: {company.name}\n"
                    f"⚠️ <b>تنبيه:</b> لم يتم العثور على منشئ الشركة لمنحه الصلاحيات تلقائياً.",
                    parse_mode="HTML",
                )
                await callback.answer("✅ تم الاعتماد (بدون ربط منشئ)", show_alert=False)
                return

            stmt_user = select(User).where(User.id == creator_id)
            result_user = await session.execute(stmt_user)
            creator = result_user.scalar_one_or_none()

            if creator:
                # Link user to company
                creator.company_id = company.id
                creator_tg_id = creator.telegram_user_id
                await session.commit()

                logger.info(
                    "Linked user to company | user_id=%s | company_id=%s",
                    creator.id,
                    company.id,
                )

                # Assign company_owner role
                await user_service.assign_role(
                    creator.id, "company_owner", granted_by=admin_internal_id
                )
                logger.info(
                    "Assigned company_owner role | user_id=%s", creator.id
                )
            else:
                logger.warning(
                    "Creator user not found in database"
                    " | creator_id=%s | company_id=%s",
                    creator_id,
                    company_id,
                    extra={"security": True},
                )

        await callback.message.edit_text(
            f"✅ <b>تم اعتماد الشركة بنجاح!</b>\n"
            f"🏢 الشركة: {company.name}\n"
            f"👤 المالك: تم ربطه ومنحه صلاحيات الإدارة.",
            parse_mode="HTML",
        )

        # 4. Notify the creator
        if creator_tg_id:
            try:
                await callback.bot.send_message(
                    chat_id=creator_tg_id,
                    text=(
                        f"🎉 <b>تم اعتماد شركتك بنجاح!</b>\n"
                        f"━━━━━━━━━━━━━━━━━━━\n\n"
                        f"🏢 الشركة: <b>{company.name}</b>\n\n"
                        f"يمكنك الآن البدء بإنشاء طلبات الشراء ودعوة موظفيك.\n"
                        f"استخدم /start للوصول إلى لوحة التحكم."
                    ),
                    parse_mode="HTML",
                )
            except Exception as e:
                logger.warning("Could not send approval notification: %s", e)

        await callback.answer("✅ تم الاعتماد!", show_alert=False)

    except Exception:
        logger.exception("Error approving company")
        await callback.answer("❌ حدث خطأ أثناء الاعتماد.", show_alert=True)


# ============================================================================
# Callback: Approve Supplier (Use Internal User ID)
# ============================================================================
@router.callback_query(F.data.startswith("sa_approve_supplier:"))
async def callback_approve_supplier(callback: CallbackQuery, roles: List[str]):
    """Handle approve supplier button click."""
    if not require_super_admin(roles):
        await callback.answer("⛔ غير مصرح لك.", show_alert=True)
        return

    supplier_id = int(callback.data.split(":", 1)[1])

    try:
        # Get Internal User ID from Telegram ID
        admin_user = await user_service.get_user_by_telegram_id(callback.from_user.id)
        if not admin_user:
            await callback.answer("❌ خطأ: لم يتم العثور على حساب المدير.", show_alert=True)
            return

        admin_internal_id = admin_user.id

        # Fetch linked user BEFORE approving
        linked_user = None
        async with get_session(silent=True) as session:
            stmt = select(User).where(User.supplier_id == supplier_id)
            result = await session.execute(stmt)
            linked_user = result.scalar_one_or_none()

        success = await supplier_service.approve_supplier(supplier_id, admin_internal_id)
        if not success:
            await callback.answer("❌ فشل الاعتماد أو المورد غير موجود.", show_alert=True)
            return

        # Fetch supplier name
        supplier = await supplier_service.get_supplier_by_id(supplier_id)
        supplier_name = supplier.name if supplier else "المورد"

        await callback.message.edit_text(
            f"✅ <b>تم اعتماد المورد بنجاح!</b>\n"
            f"🏭 المورد: {supplier_name}\n"
            f"👤 المالك: تم تفعيل حسابه ومنحه صلاحيات الإدارة.",
            parse_mode="HTML",
        )

        # Notify the linked user
        if linked_user:
            try:
                await user_service.assign_role(
                    linked_user.id, "supplier_owner", granted_by=admin_internal_id
                )
                await callback.bot.send_message(
                    chat_id=linked_user.telegram_user_id,
                    text=(
                        f"🎉 <b>تم اعتماد تسجيل المورد الخاص بك!</b>\n"
                        f"━━━━━━━━━━━━━━━━━━━\n\n"
                        f"🏭 المورد: <b>{supplier_name}</b>\n\n"
                        f"✅ أصبح حسابك كمورد نشطًا في النظام.\n"
                        f"استخدم /start للدخول إلى لوحة تحكم المورد."
                    ),
                    parse_mode="HTML",
                )
            except Exception as e:
                logger.warning("Could not notify supplier owner: %s", e)

        await callback.answer("✅ تم الاعتماد!", show_alert=False)

    except Exception:
        logger.exception("Error approving supplier")
        await callback.answer("❌ حدث خطأ أثناء الاعتماد.", show_alert=True)


# ============================================================================
# Callback: Start Company Rejection (FSM)
# ============================================================================
@router.callback_query(F.data.startswith("sa_reject_company_start:"))
async def callback_start_company_rejection(callback: CallbackQuery, state: FSMContext, roles: List[str]):
    """Start the rejection process - ask for reason."""
    if not require_super_admin(roles):
        await callback.answer("⛔ غير مصرح لك.", show_alert=True)
        return

    company_id = callback.data.split(":", 1)[1]
    await state.update_data(rejecting_company_id=int(company_id))

    await callback.message.answer(
        "❌ <b>رفض طلب تسجيل الشركة</b>\n"
        "━━━━━━━━━━━━━━━━━━━\n\n"
        "📝 الرجاء إدخال سبب الرفض:",
        parse_mode="HTML",
    )
    await state.set_state(OwnerRejectFlow.waiting_for_company_rejection_reason)
    await callback.answer()


# ============================================================================
# Callback: Start Supplier Rejection (FSM)
# ============================================================================
@router.callback_query(F.data.startswith("sa_reject_supplier_start:"))
async def callback_start_supplier_rejection(callback: CallbackQuery, state: FSMContext, roles: List[str]):
    """Start the supplier rejection process - ask for reason."""
    if not require_super_admin(roles):
        await callback.answer("⛔ غير مصرح لك.", show_alert=True)
        return

    supplier_id = callback.data.split(":", 1)[1]
    await state.update_data(rejecting_supplier_id=int(supplier_id))

    await callback.message.answer(
        "❌ <b>رفض طلب تسجيل المورد</b>\n"
        "━━━━━━━━━━━━━━━━━━━\n\n"
        "📝 الرجاء إدخال سبب الرفض:",
        parse_mode="HTML",
    )
    await state.set_state(OwnerRejectFlow.waiting_for_supplier_rejection_reason)
    await callback.answer()


# ============================================================================
# FSM: Receive Company Rejection Reason
# ============================================================================
@router.message(OwnerRejectFlow.waiting_for_company_rejection_reason)
async def process_company_rejection_reason(message: Message, state: FSMContext, roles: List[str]):
    """Process company rejection reason input."""
    if not require_super_admin(roles):
        await message.answer("⛔ غير مصرح لك.")
        await state.clear()
        return

    reason = message.text.strip()
    data = await state.get_data()
    company_id = data.get("rejecting_company_id")

    if not company_id or not reason:
        await message.answer("❌ حدث خطأ. الرجاء البدء من جديد.")
        await state.clear()
        return

    admin_user = await user_service.get_user_by_telegram_id(message.from_user.id)
    if not admin_user:
        await message.answer("❌ خطأ: لم يتم العثور على حساب المدير.")
        await state.clear()
        return

    admin_internal_id = admin_user.id

    success = await company_service.reject_company(
        company_id, admin_internal_id, reason=reason
    )

    if success:
        await message.answer(
            f"❌ <b>تم رفض الشركة بنجاح.</b>\n"
            f"📝 السبب: {reason}",
            parse_mode="HTML",
        )
    else:
        await message.answer("❌ فشل الرفض.")

    await state.clear()


# ============================================================================
# FSM: Receive Supplier Rejection Reason
# ============================================================================
@router.message(OwnerRejectFlow.waiting_for_supplier_rejection_reason)
async def process_supplier_rejection_reason(message: Message, state: FSMContext, roles: List[str]):
    """Process supplier rejection reason input."""
    if not require_super_admin(roles):
        await message.answer("⛔ غير مصرح لك.")
        await state.clear()
        return

    reason = message.text.strip()
    data = await state.get_data()
    supplier_id = data.get("rejecting_supplier_id")

    if not supplier_id or not reason:
        await message.answer("❌ حدث خطأ. الرجاء البدء من جديد.")
        await state.clear()
        return

    admin_user = await user_service.get_user_by_telegram_id(message.from_user.id)
    if not admin_user:
        await message.answer("❌ خطأ: لم يتم العثور على حساب المدير.")
        await state.clear()
        return

    admin_internal_id = admin_user.id

    # ✅ Pass reason so it is recorded in audit_log
    success = await supplier_service.reject_supplier(
        supplier_id, admin_internal_id, reason=reason
    )

    if success:
        await message.answer(
            f"❌ <b>تم رفض المورد بنجاح.</b>\n"
            f"📝 السبب: {reason}",
            parse_mode="HTML",
        )
    else:
        await message.answer("❌ فشل الرفض.")

    await state.clear()


# ============================================================================
# Callback: Export report to Excel
# ============================================================================
@router.callback_query(F.data == "sa_export_report_excel")
async def callback_export_report(callback: CallbackQuery, roles: List[str]):
    """Handle export report to Excel button."""
    if not require_super_admin(roles):
        await callback.answer("⛔ غير مصرح لك.", show_alert=True)
        return

    await callback.answer("⏳ جاري إنشاء التقرير...", show_alert=False)

    try:
        async with get_session(silent=True) as session:
            # Fetch Companies
            companies = (await session.execute(
                select(Company)
                .where(Company.status == CompanyStatus.active)
                .order_by(Company.created_at.desc())
            )).scalars().all()

            companies_data = []
            for c in companies:
                companies_data.append({
                    "name": c.name,
                    "inn": c.inn or "N/A",
                    "city": c.city or "N/A",
                    "created_at": c.created_at.strftime("%Y-%m-%d %H:%M") if c.created_at else "N/A",
                })

            # Fetch Users
            users = (await session.execute(
                select(User, Company.name)
                .outerjoin(Company, User.company_id == Company.id)
                .where(User.status == UserStatus.active)
                .order_by(User.created_at.desc())
            )).all()

            users_data = []
            for u, company_name in users:
                users_data.append({
                    "full_name": u.full_name,
                    "telegram_id": u.telegram_user_id,
                    "company": company_name or "N/A",
                    "created_at": u.created_at.strftime("%Y-%m-%d %H:%M") if u.created_at else "N/A",
                })

        # Generate Excel
        wb = Workbook()

        # Sheet 1: Companies
        ws_c = wb.active
        ws_c.title = "الشركات"
        ws_c.append(["اسم الشركة", "الرقم الضريبي (INN)", "المدينة", "تاريخ التسجيل"])
        for c in companies_data:
            ws_c.append([c["name"], c["inn"], c["city"], c["created_at"]])

        # Sheet 2: Users
        ws_u = wb.create_sheet("المستخدمون")
        ws_u.append(["الاسم الكامل", "Telegram ID", "الشركة", "تاريخ التسجيل"])
        for u in users_data:
            ws_u.append([u["full_name"], u["telegram_id"], u["company"], u["created_at"]])

        # Auto-width
        for ws in [ws_c, ws_u]:
            for column in ws.columns:
                max_length = max(len(str(cell.value or "")) for cell in column)
                ws.column_dimensions[get_column_letter(column[0].column)].width = min(max_length + 2, 50)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"procurement_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

        await callback.message.answer_document(
            document=BufferedInputFile(output.getvalue(), filename=filename),
            caption=f"📊 <b>تقرير Procurement Bot</b>\n📅 {datetime.now().strftime('%Y-%m-%d %H:%M')}",
            parse_mode="HTML",
        )

        await callback.answer("✅ تم تصدير التقرير بنجاح!")

    except Exception:
        logger.exception("Error exporting report")
        await callback.message.answer("❌ عذراً، حدث خطأ أثناء تصدير التقرير.")
        await callback.answer("❌ حدث خطأ", show_alert=True)